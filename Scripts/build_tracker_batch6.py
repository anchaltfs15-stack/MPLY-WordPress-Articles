import openpyxl
from openpyxl.styles import Font, Alignment

BRIEF = r"C:\Users\SrimanNarayan\OneDrive - TransForm Solutions (P) Limited\MFunds_WP Articles\Production Briefs\2026-09-28_Set_01.xlsx"
TRACKER = r"C:\Users\SrimanNarayan\OneDrive - TransForm Solutions (P) Limited\MFunds_WP Articles\Production Trackers\MPLY_Article_Production.xlsx"

wb_brief = openpyxl.load_workbook(BRIEF)
ws_brief = wb_brief["Set 01"]
brief_rows = {}
for r in range(2, ws_brief.max_row + 1):
    no = ws_brief.cell(row=r, column=1).value
    brief_rows[no] = {
        "title": ws_brief.cell(row=r, column=2).value,
        "primary_kw": ws_brief.cell(row=r, column=3).value,
        "secondary_kw": ws_brief.cell(row=r, column=4).value,
        "related": ws_brief.cell(row=r, column=5).value,
        "cluster": ws_brief.cell(row=r, column=6).value,
        "intent": ws_brief.cell(row=r, column=7).value,
        "relevance": ws_brief.cell(row=r, column=8).value,
        "connection_logic": ws_brief.cell(row=r, column=9).value,
        "placement_brief": ws_brief.cell(row=r, column=10).value,
    }

PROD = [
    dict(no=44, seo="What Is Earnings Quality",
         meta="Earnings quality measures how much of reported profit is backed by real cash versus accounting estimates. The accrual gap and where weak quality shows up.",
         mply_decision="Contextual Section (revised 2026-09-24: Omit retired, comprehensive MPLY now mandatory)", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative net-income-vs-operating-cash-flow line chart (accrual gap concept), labeled illustrative.",
         internal_links="", ext_sources="strategysharesetfs.com/mply",
         hashtags="#EarningsQuality #QualityInvesting #Accounting #FundamentalAnalysis #MPLY #Investing",
         filename="044_What_Is_Earnings_Quality.docx", qc="PASS", notes="Explainer archetype. MPLY angle: earnings quality as a complementary diligence check on MPLY holdings' reported margins."),
    dict(no=45, seo="What Is a Cash Conversion Cycle",
         meta="The cash conversion cycle measures how fast a business gets its cash back. The DIO+DSO-DPO formula and what a negative cycle means.",
         mply_decision="Contextual Section (revised 2026-09-24)", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative DIO+DSO-DPO cash conversion cycle breakdown chart, labeled illustrative.",
         internal_links="", ext_sources="strategysharesetfs.com/mply",
         hashtags="#CashConversionCycle #WorkingCapital #BusinessEconomics #MPLY #Investing #Fundamentals",
         filename="045_What_Is_a_Cash_Conversion_Cycle.docx", qc="PASS", notes="Explainer archetype. MPLY angle: favorable cash conversion cycles as corroborating evidence of pricing power/industry concentration."),
    dict(no=46, seo="What Is Operating Leverage",
         meta="Operating leverage describes how much profit swings relative to revenue based on fixed versus variable cost mix. Why it amplifies both gains and losses.",
         mply_decision="Contextual Section (revised 2026-09-24)", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative profit-swing-by-cost-structure bar chart, labeled illustrative.",
         internal_links="", ext_sources="strategysharesetfs.com/mply",
         hashtags="#OperatingLeverage #BusinessEconomics #Profitability #MPLY #Investing #Fundamentals",
         filename="046_What_Is_Operating_Leverage.docx", qc="PASS", notes="Mechanism-Based archetype. MPLY angle: dominant/scaled businesses often carry high operating leverage, cuts both ways as a risk factor."),
    dict(no=47, seo="What Is Gross Margin and Why It Matters",
         meta="Gross margin shows how much of each sales dollar survives direct production costs. The formula and why the trend matters more than one reading.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative revenue-to-operating-profit margin waterfall chart (shared with No.48), labeled illustrative.",
         internal_links="Anchor 'pricing power' -> What Is Pricing Power (No.1)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#GrossMargin #QualityInvesting #PricingPower #MPLY #Investing #Fundamentals",
         filename="047_What_Is_Gross_Margin_and_Why_It_Matters.docx", qc="PASS", notes="Explainer archetype."),
    dict(no=48, seo="What Is Operating Margin",
         meta="Operating margin captures profit after both production costs and operating expenses. The gap between gross and operating margin, and what it reveals.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative revenue-to-operating-profit margin waterfall chart (shared with No.47), labeled illustrative.",
         internal_links="Anchor 'gross margin' -> What Is Gross Margin and Why It Matters (No.47)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#OperatingMargin #QualityInvesting #Profitability #MPLY #Investing #Fundamentals",
         filename="048_What_Is_Operating_Margin.docx", qc="PASS", notes="Explainer archetype."),
    dict(no=49, seo="What Is a Recurring Revenue Business",
         meta="A recurring revenue business is structured around repeated payments, but predictability and durability are separate questions. Where the moat question begins.",
         mply_decision="Contextual Section, expanded per Strong-thematic rule", mply_placement="Contextual Section (expanded)",
         visual="Yes", visual_brief="Illustrative churn-vs-lifetime-value sensitivity chart (shared with No.50, No.51), labeled illustrative.",
         internal_links="Anchor 'recurring revenue moat' -> How Recurring Revenue Creates a Business Moat (No.9)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#RecurringRevenue #BusinessModels #SubscriptionBusiness #EconomicMoat #MPLY #Investing",
         filename="049_What_Is_a_Recurring_Revenue_Business.docx", qc="PASS", notes="Explainer archetype; business-model-category framing distinct from No.9's moat-mechanism analysis."),
    dict(no=50, seo="What Is a Subscription Business Model",
         meta="A subscription business model is one implementation of recurring revenue, defined by its own metrics: CAC, churn, and lifetime value.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative churn-vs-lifetime-value sensitivity chart (shared with No.49, No.51), labeled illustrative.",
         internal_links="Anchor 'recurring revenue business' -> What Is a Recurring Revenue Business (No.49)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#SubscriptionBusiness #BusinessModels #RecurringRevenue #Churn #MPLY #Investing",
         filename="050_What_Is_Subscription_Business_Model.docx", qc="PASS", notes="Comparison archetype; operating-mechanics angle distinct from No.49's category definition."),
    dict(no=51, seo="What Is Customer Lifetime Value",
         meta="Customer lifetime value estimates total expected revenue from a customer relationship. The formula and why it's especially sensitive to churn.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative churn-vs-lifetime-value sensitivity chart (shared with No.49, No.50), labeled illustrative.",
         internal_links="Anchor 'CAC vs LTV' -> Customer Acquisition Cost vs Lifetime Value (No.52)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#CustomerLifetimeValue #BusinessEconomics #Retention #MPLY #Investing #Fundamentals",
         filename="051_What_Is_Customer_Lifetime_Value.docx", qc="PASS", notes="Explainer archetype."),
    dict(no=52, seo="Customer Acquisition Cost vs Lifetime Value",
         meta="The LTV:CAC ratio and payback period together determine whether a growth strategy is actually economical, not either metric alone.",
         mply_decision="Contextual Section (revised 2026-09-24)", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Illustrative LTV:CAC ratio benchmark chart, labeled illustrative.",
         internal_links="Anchor 'customer lifetime value' -> What Is Customer Lifetime Value (No.51)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#CAC #CustomerLifetimeValue #UnitEconomics #BusinessEconomics #MPLY #Investing",
         filename="052_Customer_Acquisition_Cost_vs_Lifetime_Value.docx", qc="PASS", notes="Comparison archetype. MPLY angle: dominant companies typically show stronger LTV:CAC ratios via lower paid-acquisition dependence."),
    dict(no=53, seo="What Is a Monopoly Business? The De Beers Case",
         meta="De Beers controlled up to 90% of world diamond production for decades. How the monopoly formed, and how it eroded despite near-total market control.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="De Beers market share decline bar chart (peak vs. early-2000s), verified data.",
         internal_links="Anchor 'market dominance vs monopoly' -> Market Dominance vs Monopoly (No.28)",
         ext_sources="en.wikipedia.org/wiki/De_Beers; strategysharesetfs.com/mply",
         hashtags="#Monopoly #MarketDominance #MPLY #EconomicMoat #CompetitiveAdvantage #Investing",
         filename="053_What_Is_a_Monopoly_Business.docx", qc="PASS", notes="Case Study/Examples archetype; De Beers angle distinct from No.28's legal/economic comparison framing. FAQ included."),
]

wb_t = openpyxl.load_workbook(TRACKER)
ws = wb_t["Production Tracker"]
start_row = ws.max_row + 1

BODY_FONT = Font(name="Arial")
WRAP = Alignment(wrap_text=True, vertical="top")

for i, item in enumerate(PROD):
    r = start_row + i
    b = brief_rows[item["no"]]
    row_vals = [
        "2026-09-28", "Set 01", item["no"], b["title"], b["primary_kw"], b["secondary_kw"],
        b["related"], b["cluster"], b["intent"], b["relevance"], b["connection_logic"],
        item["seo"], item["meta"], item["mply_decision"], item["mply_placement"],
        item["visual"], item["visual_brief"], item["internal_links"], item["ext_sources"],
        item["hashtags"], item["filename"], item["qc"], item["notes"],
    ]
    for c, v in enumerate(row_vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = BODY_FONT
        cell.alignment = WRAP
    ws.row_dimensions[r].height = 30

wb_t.save(TRACKER)
print(f"Appended {len(PROD)} rows starting at row {start_row} to {TRACKER}")
