import openpyxl
from openpyxl.styles import Font, Alignment

BRIEF = r"C:\Users\SrimanNarayan\OneDrive - TransForm Solutions (P) Limited\MFunds_WP Articles\Production Briefs\2026-09-28_Set_01.xlsx"
TRACKER = r"C:\Users\SrimanNarayan\OneDrive - TransForm Solutions (P) Limited\MFunds_WP Articles\Production Trackers\MPLY_Article_Production.xlsx"

wb_brief = openpyxl.load_workbook(BRIEF)
ws_brief = wb_brief["Set 01"]
brief_rows = {}
for r in range(2, ws_brief.max_row + 1):
    no = ws_brief.cell(row=r, column=1).value
    brief_rows[no] = {
        "title": ws_brief.cell(row=r, column=2).value,
        "primary_kw": ws_brief.cell(row=r, column=3).value,
        "secondary_kw": ws_brief.cell(row=r, column=4).value,
        "related": ws_brief.cell(row=r, column=5).value,
        "cluster": ws_brief.cell(row=r, column=6).value,
        "intent": ws_brief.cell(row=r, column=7).value,
        "relevance": ws_brief.cell(row=r, column=8).value,
        "connection_logic": ws_brief.cell(row=r, column=9).value,
        "placement_brief": ws_brief.cell(row=r, column=10).value,
    }

PROD = [
    dict(no=12, seo="What Is a Regulatory Moat? Credit Rating Agencies Explained",
         meta="A regulatory moat exists when government, not the market, limits competition. Credit rating agencies show how a 1975 SEC designation still concentrates an industry.",
         mply_decision="Contextual Section", mply_placement="Brief Mention (per brief; expanded per Strong-thematic depth rule)",
         visual="Yes", visual_brief="Credit rating revenue concentration donut chart (S&P/Moody's/Fitch 91%), verified data.",
         internal_links="Anchor 'barriers to entry' -> What Are Barriers to Entry in Business (No.6)",
         ext_sources="congress.gov/crs-product/IF11916; strategysharesetfs.com/mply",
         hashtags="#RegulatoryMoat #EconomicMoat #CompetitiveAdvantage #BarriersToEntry #MPLY #Investing",
         filename="012_What_Is_a_Regulatory_Moat.docx", qc="PASS", notes="Explainer archetype."),
    dict(no=14, seo="Economies of Scale vs Competitive Advantage",
         meta="Scale lowers costs, but it doesn't always create a moat. Amazon's fulfillment network versus commodity steel production shows the difference.",
         mply_decision="Brief Mention", mply_placement="Brief Mention",
         visual="No", visual_brief="Evaluated; economies-of-scale chart already exists for article 013, no new verified data warranted a second chart.",
         internal_links="Anchor 'cost advantage moat' -> What Is a Cost Advantage Moat (No.13)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#EconomiesOfScale #CompetitiveAdvantage #EconomicMoat #MPLY #Investing #Fundamentals",
         filename="014_Economies_of_Scale_vs_Competitive_Advantage.docx", qc="PASS", notes="Comparison archetype."),
    dict(no=15, seo="Economies of Scope vs Economies of Scale",
         meta="Scope and scale both lower costs but through different mechanisms. P&G's shared resources across categories illustrate scope; a single-product plant illustrates scale.",
         mply_decision="Brief Mention", mply_placement="Brief Mention",
         visual="No", visual_brief="Evaluated; checklist-style bullets in-text served the comparison without a separate chart.",
         internal_links="",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#EconomiesOfScope #EconomiesOfScale #BusinessEconomics #CompetitiveAdvantage #MPLY #Investing",
         filename="015_Economies_of_Scope_vs_Economies_of_Scale.docx", qc="PASS", notes="Checklist/Research Guide archetype; Business Economics cluster (not Economic Moats)."),
    dict(no=16, seo="What Is a Learning Curve Advantage",
         meta="A learning curve advantage comes from accumulated production experience, not current volume. Wright's 1936 aircraft study and solar panel costs show the pattern.",
         mply_decision="Brief Mention", mply_placement="Brief Mention",
         visual="Yes", visual_brief="Illustrative learning-curve (Wright's Law shape) chart.",
         internal_links="Anchor 'economies of scale' -> Economies of Scale vs Competitive Advantage (No.14)",
         ext_sources="ourworldindata.org/learning-curve; en.wikipedia.org/wiki/Experience_curve_effect",
         hashtags="#LearningCurve #ExperienceCurve #CompetitiveAdvantage #EconomicMoat #Investing #Fundamentals",
         filename="016_What_Is_a_Learning_Curve_Advantage.docx", qc="PASS", notes="Mechanism-Based archetype."),
    dict(no=17, seo="What Is a Distribution Moat",
         meta="A distribution moat comes from customer access a competitor can't easily replicate. Coca-Cola's global bottling network shows how it works and where it's weakening.",
         mply_decision="Contextual Section", mply_placement="Brief Mention (per brief; expanded per Strong-thematic depth rule)",
         visual="No", visual_brief="Evaluated; no verified numeric dataset available without fabricating distribution-reach figures.",
         internal_links="Anchor 'economic moat examples' -> Economic Moat Examples (No.3)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#DistributionMoat #EconomicMoat #CompetitiveAdvantage #SupplyChain #MPLY #Investing",
         filename="017_What_Is_a_Distribution_Moat.docx", qc="PASS", notes="Case Study/Examples archetype; Coca-Cola distribution angle distinct from article 003's brand angle."),
    dict(no=18, seo="What Is an Ecosystem Moat? Apple's Lock-In Explained",
         meta="An ecosystem moat comes from multiple connected products that work better together. Apple's hardware, software, and services lock-in shows how it compounds.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Hub-and-spoke ecosystem conceptual diagram.",
         internal_links="Anchor 'switching costs' -> What Are Switching Costs and Why Do They Matter (No.4)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#EcosystemMoat #EconomicMoat #CompetitiveAdvantage #SwitchingCosts #MPLY #Investing #Fundamentals",
         filename="018_What_Is_an_Ecosystem_Moat.docx", qc="PASS", notes="Explainer archetype; FAQ included (broad, high-intent topic)."),
    dict(no=19, seo="What Is a Data Moat",
         meta="A data moat exists when proprietary data creates a feedback loop that improves a product. Google's search quality shows how it works, and when it doesn't.",
         mply_decision="Brief Mention", mply_placement="Brief Mention",
         visual="No", visual_brief="Evaluated; no verified numeric dataset available without fabricating data-scale figures.",
         internal_links="Anchor 'network effects' -> Network Effects Explained (No.5)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#DataMoat #CompetitiveAdvantage #EconomicMoat #MachineLearning #MPLY #Investing",
         filename="019_What_Is_a_Data_Moat.docx", qc="PASS", notes="Investor Framework archetype."),
    dict(no=20, seo="What Is an Intellectual Property Moat",
         meta="An IP moat is a legally protected advantage with a known expiration date. Here's how patents work, the pharmaceutical patent cliff, and trade secrets.",
         mply_decision="Contextual Section", mply_placement="Brief Mention (per brief; expanded per Strong-thematic depth rule)",
         visual="No", visual_brief="Evaluated; a patent-lifecycle timeline was considered but the prose Timeline/Evolution structure already covers this without a chart.",
         internal_links="Anchor 'barriers to entry' -> What Are Barriers to Entry in Business (No.6)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#IntellectualProperty #PatentMoat #EconomicMoat #CompetitiveAdvantage #MPLY #Investing",
         filename="020_What_Is_an_Intellectual_Property_Moat.docx", qc="PASS", notes="Timeline/Evolution archetype; FAQ included (broad, high-intent topic)."),
    dict(no=21, seo="What Is a Scale Moat? Amazon's Flywheel Explained",
         meta="A scale moat is self-reinforcing, not just cost-reducing. Amazon's flywheel shows how scale can compound rather than just lower costs.",
         mply_decision="Brief Mention", mply_placement="Brief Mention",
         visual="No", visual_brief="Evaluated; the flywheel concept is well conveyed in prose; a circular-diagram chart was considered but judged unnecessary given No.18's ecosystem diagram already fills that visual role in this batch.",
         internal_links="Anchor 'economies of scale' -> Economies of Scale vs Competitive Advantage (No.14)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#ScaleMoat #EconomicMoat #CompetitiveAdvantage #MPLY #Investing #Fundamentals",
         filename="021_What_Is_a_Scale_Moat.docx", qc="PASS", notes="Mechanism-Based archetype."),
    dict(no=22, seo="What Is a Market Share Moat? Lessons From Nokia",
         meta="Dominant market share looks like a moat but isn't automatically one. Nokia's collapse from 51% to under 3% share shows why the underlying driver matters more.",
         mply_decision="Contextual Section", mply_placement="Brief Mention (per brief; expanded per Strong-thematic depth rule)",
         visual="Yes", visual_brief="Nokia 2007 vs 2013 market share collapse bar chart, verified data.",
         internal_links="Anchor 'second mover advantage' -> What Is a Second Mover Advantage (No.33)",
         ext_sources="en.wikipedia.org/wiki/The_Decline_and_Fall_of_Nokia; strategysharesetfs.com/mply",
         hashtags="#MarketShare #EconomicMoat #CompetitiveAdvantage #MPLY #Investing #Fundamentals",
         filename="022_What_Is_a_Market_Share_Moat.docx", qc="PASS", notes="Myth vs. Reality archetype."),
]

wb_t = openpyxl.load_workbook(TRACKER)
ws = wb_t["Production Tracker"]
start_row = ws.max_row + 1

BODY_FONT = Font(name="Arial")
WRAP = Alignment(wrap_text=True, vertical="top")

for i, item in enumerate(PROD):
    r = start_row + i
    b = brief_rows[item["no"]]
    row_vals = [
        "2026-09-28", "Set 01", item["no"], b["title"], b["primary_kw"], b["secondary_kw"],
        b["related"], b["cluster"], b["intent"], b["relevance"], b["connection_logic"],
        item["seo"], item["meta"], item["mply_decision"], item["mply_placement"],
        item["visual"], item["visual_brief"], item["internal_links"], item["ext_sources"],
        item["hashtags"], item["filename"], item["qc"], item["notes"],
    ]
    for c, v in enumerate(row_vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = BODY_FONT
        cell.alignment = WRAP
    ws.row_dimensions[r].height = 30

wb_t.save(TRACKER)
print(f"Appended {len(PROD)} rows starting at row {start_row} to {TRACKER}")
