import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

BRIEF = r"C:\Users\SrimanNarayan\OneDrive - TransForm Solutions (P) Limited\MFunds_WP Articles\Production Briefs\2026-09-28_Set_01.xlsx"
TRACKER = r"C:\Users\SrimanNarayan\OneDrive - TransForm Solutions (P) Limited\MFunds_WP Articles\Production Trackers\MPLY_Article_Production.xlsx"

wb_brief = openpyxl.load_workbook(BRIEF)
ws_brief = wb_brief["Set 01"]
brief_rows = {}
for r in range(2, ws_brief.max_row + 1):
    no = ws_brief.cell(row=r, column=1).value
    brief_rows[no] = {
        "title": ws_brief.cell(row=r, column=2).value,
        "primary_kw": ws_brief.cell(row=r, column=3).value,
        "secondary_kw": ws_brief.cell(row=r, column=4).value,
        "related": ws_brief.cell(row=r, column=5).value,
        "cluster": ws_brief.cell(row=r, column=6).value,
        "intent": ws_brief.cell(row=r, column=7).value,
        "relevance": ws_brief.cell(row=r, column=8).value,
        "connection_logic": ws_brief.cell(row=r, column=9).value,
        "placement_brief": ws_brief.cell(row=r, column=10).value,
    }

# Per-article production data for the 14 completed articles.
PROD = [
    dict(no=1, seo="What Is Pricing Power? Why It Matters for Investors",
         meta="Pricing power lets a company raise prices without losing customers. Here's how it works, how to spot it, and why MPLY treats it as a core screening signal.",
         mply_decision="Brief Mention section, expanded to substantial", mply_placement="Contextual Section (upgraded from brief's Brief Mention)",
         visual="No", visual_brief="Evaluated; no verified numeric dataset fit the topic without fabricating data.",
         internal_links="Anchor 'economic moat' -> What Is a Competitive Moat in Investing (No.2); Anchor 'switching costs' -> What Are Switching Costs and Why Do They Matter (No.4)",
         ext_sources="strategysharesetfs.com/mply; stockanalysis.com/etf/mply",
         hashtags="#PricingPower #CompetitiveAdvantage #EconomicMoat #MPLY #Investing #QualityInvesting #Fundamentals",
         filename="001_What_Is_Pricing_Power.docx", qc="PASS", notes="Pilot article. MPLY section expanded from 1 sentence to a full section after user feedback.")
    ,
    dict(no=2, seo="What Is a Competitive Moat in Investing",
         meta="A competitive moat is a durable advantage that protects a business's profits from competition. Here's where the term comes from, its common sources, and its limits.",
         mply_decision="Contextual Section (substantial, foundational)", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Six Sources of a Durable Competitive Advantage conceptual framework diagram (built).",
         internal_links="Anchor 'economic moat examples' -> Economic Moat Examples (No.3); Anchor 'switching costs' -> What Are Switching Costs (No.4)",
         ext_sources="en.wikipedia.org/wiki/Economic_moat; inc.com (Blockbuster/Kodak/Xerox failure article); strategysharesetfs.com/mply",
         hashtags="#CompetitiveMoat #EconomicMoat #CompetitiveAdvantage #MPLY #ValueInvesting #Investing #Fundamentals",
         filename="002_What_Is_a_Competitive_Moat.docx", qc="PASS", notes="Umbrella/foundational article for the Economic Moats cluster.")
    ,
    dict(no=3, seo="Economic Moat Examples: Coca-Cola, Visa, and ASML",
         meta="Real examples of durable competitive advantages: Coca-Cola's brand, Visa and Mastercard's network effects, and ASML's technology barrier to entry.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="No", visual_brief="Evaluated; a comparison table was considered but the prose already covers the 3 examples clearly without one.",
         internal_links="Anchor 'competitive moat' -> What Is a Competitive Moat in Investing (No.2)",
         ext_sources="en.wikipedia.org/wiki/Economic_moat; strategysharesetfs.com/mply",
         hashtags="#EconomicMoat #CompetitiveAdvantage #BrandMoat #NetworkEffects #MPLY #Investing #ValueInvesting",
         filename="003_Economic_Moat_Examples.docx", qc="PASS", notes="Named, researched examples per specificity rule (no vague 'a dominant company emerged' language).")
    ,
    dict(no=4, seo="What Are Switching Costs and Why Do They Matter",
         meta="Switching costs make it expensive or disruptive for a customer to leave, even when a competitor's product is better. Here's the mechanism, evidence, and limits.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="No", visual_brief="Evaluated; no verified numeric dataset available without fabricating churn/retention figures.",
         internal_links="Anchor 'customer retention' -> Customer Retention as a Competitive Advantage (No.10)",
         ext_sources="fcc.gov (mobile number portability); strategysharesetfs.com/mply",
         hashtags="#SwitchingCosts #EconomicMoat #CompetitiveAdvantage #CustomerRetention #Investing #Fundamentals",
         filename="004_What_Are_Switching_Costs.docx", qc="PASS", notes="")
    ,
    dict(no=5, seo="Network Effects Explained: A Competitive Advantage",
         meta="A network effect makes a product more valuable as more people use it. Here's the mechanism, its origin, and why it doesn't guarantee durability.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="How Network Effects Compound conceptual illustrative curve (built).",
         internal_links="Anchor 'network effect moat' -> What Is a Network Effect Moat (No.11)",
         ext_sources="en.wikipedia.org/wiki/Network_effect; strategysharesetfs.com/mply",
         hashtags="#NetworkEffects #EconomicMoat #CompetitiveAdvantage #PlatformEconomics #MPLY #Investing",
         filename="005_Network_Effects_Explained.docx", qc="PASS", notes="Definitional companion to No.11's durability-focused case study; different angle/archetype (Timeline/Evolution vs. Case Study).")
    ,
    dict(no=6, seo="What Are Barriers to Entry in Business",
         meta="Barriers to entry make it difficult or expensive for a new competitor to enter a market. Here's how capital, regulation, and scale each create one.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="No", visual_brief="Evaluated; a checklist-style visual was considered but the bulleted questions in-text already served the purpose.",
         internal_links="Anchor 'competitive moat' -> What Is a Competitive Moat in Investing (No.2)",
         ext_sources="fcc.gov/auctions (spectrum licensing); strategysharesetfs.com/mply",
         hashtags="#BarriersToEntry #EconomicMoat #CompetitiveAdvantage #MarketStructure #MPLY #Investing",
         filename="006_What_Are_Barriers_to_Entry.docx", qc="PASS", notes="")
    ,
    dict(no=7, seo="Brand Loyalty as a Competitive Advantage",
         meta="Brand loyalty can support pricing power and retention, but it isn't automatic or permanent. Here's the myth, the reality, and Harley-Davidson's example.",
         mply_decision="Brief Mention, expanded to two paragraphs", mply_placement="Brief Mention (per brief; expanded per Strong-thematic depth rule)",
         visual="No", visual_brief="Evaluated; no verified numeric dataset (e.g. repeat-purchase rates) available for this topic without fabricating data.",
         internal_links="Anchor 'pricing power' -> What Is Pricing Power (No.1)",
         ext_sources="en.wikipedia.org/wiki/Harley-Davidson; strategysharesetfs.com/mply",
         hashtags="#BrandLoyalty #CompetitiveAdvantage #EconomicMoat #PricingPower #Investing #Fundamentals",
         filename="007_Brand_Loyalty_Competitive_Advantage.docx", qc="PASS", notes="Myth vs. Reality archetype.")
    ,
    dict(no=8, seo="What Makes a Business Hard to Displace",
         meta="The businesses hardest to displace combine multiple reinforcing advantages, not just one. Here's the framework for evaluating layered competitive advantages.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="No", visual_brief="Evaluated; a framework diagram was considered but the bulleted indicators list already serves this purpose without duplicating No.2's chart.",
         internal_links="Anchor 'economic moat examples' -> Economic Moat Examples (No.3)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#CompetitiveAdvantage #EconomicMoat #BusinessDefensibility #MPLY #Investing #Fundamentals",
         filename="008_What_Makes_a_Business_Hard_to_Displace.docx", qc="PASS", notes="Investor Framework archetype; synthesizes the cluster's individual moat-source articles.")
    ,
    dict(no=9, seo="How Recurring Revenue Creates a Business Moat",
         meta="Recurring revenue and a durable moat aren't the same thing. Here's the difference, illustrated by a gym membership versus Adobe Creative Cloud.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="No", visual_brief="Evaluated; no verified numeric dataset (e.g. actual churn rates) available without fabricating data.",
         internal_links="Anchor 'switching costs' -> What Are Switching Costs and Why Do They Matter (No.4)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#RecurringRevenue #EconomicMoat #CompetitiveAdvantage #SubscriptionBusiness #MPLY #Investing",
         filename="009_Recurring_Revenue_Business_Moat.docx", qc="PASS", notes="Comparison archetype (recurring revenue vs. genuine durability).")
    ,
    dict(no=10, seo="Customer Retention as a Competitive Advantage",
         meta="High customer retention looks like a moat, but the underlying driver matters more than the number. Here's how to tell structural retention from inertia.",
         mply_decision="Brief Mention", mply_placement="Brief Mention",
         visual="No", visual_brief="Evaluated; no verified numeric dataset available without fabricating churn/retention data.",
         internal_links="Anchor 'switching costs' -> What Are Switching Costs and Why Do They Matter (No.4)",
         ext_sources="strategysharesetfs.com/mply",
         hashtags="#CustomerRetention #CompetitiveAdvantage #EconomicMoat #ChurnRate #MPLY #Investing",
         filename="010_Customer_Retention_Competitive_Advantage.docx", qc="PASS", notes="Brief kept intentionally lighter: brief's own Connection Logic frames retention as an indirect signal, consistent with the standing MPLY-tag rule's exception.")
    ,
    dict(no=11, seo="What Is a Network Effect Moat? Lessons From Friendster",
         meta="A network effect moat only holds if it survives real competitive pressure. Friendster's rise and fall versus Visa's durable network effect shows why.",
         mply_decision="Contextual Section", mply_placement="Contextual Section",
         visual="Yes", visual_brief="Friendster lifecycle timeline chart, verified dates (built).",
         internal_links="Anchor 'network effects' -> Network Effects Explained (No.5)",
         ext_sources="en.wikipedia.org/wiki/Friendster; technologyreview.com (MIT Tech Review Friendster autopsy); strategysharesetfs.com/mply",
         hashtags="#NetworkEffects #EconomicMoat #CompetitiveAdvantage #PlatformEconomics #MPLY #Investing",
         filename="011_What_Is_a_Network_Effect_Moat.docx", qc="PASS", notes="Case Study/Examples archetype, durability-focused angle distinct from No.5's definitional treatment.")
    ,
    dict(no=13, seo="What Is a Cost Advantage Moat in Investing?",
         meta="A cost advantage moat lets a company underprice competitors and still profit. Here's what to look for, real examples, and how MPLY's screen connects.",
         mply_decision="Contextual Section, expanded from throwaway clause", mply_placement="Brief Mention (per brief; expanded per adequate-coverage standard)",
         visual="Yes", visual_brief="Economies-of-scale illustrative conceptual chart (built).",
         internal_links="Anchor 'economies of scale' -> Economic Moat Examples (No.3)",
         ext_sources="simpleflying.com (Southwest Airlines fleet strategy); strategysharesetfs.com/mply",
         hashtags="#CostAdvantageMoat #EconomicMoat #CompetitiveAdvantage #MPLY #Investing #ValueInvesting #Fundamentals",
         filename="013_What_Is_a_Cost_Advantage_Moat.docx", qc="PASS", notes="Pilot article, filename corrected from 002 to 013 (brief row 13) to avoid collision with new batch. MPLY paragraph rewritten from a throwaway clause after user feedback.")
    ,
    dict(no=33, seo="What Is a Second Mover Advantage in Business?",
         meta="Being first to market isn't always the winning move. AltaVista, Myspace, and early touchscreen phones show how second movers can win.",
         mply_decision="Brief Mention, added after initial Omit", mply_placement="Omit (per brief; overridden after review found a genuine connection)",
         visual="Yes", visual_brief="Pioneer-vs-second-mover timeline chart (built).",
         internal_links="Anchor 'network effects' -> Network Effects Explained (No.5); Anchor 'switching costs' -> What Are Switching Costs (No.4)",
         ext_sources="cybercultural.com; en.wikipedia.org/wiki/AltaVista; pcworld.com; appleinsider.com; simpletexting.com",
         hashtags="#SecondMoverAdvantage #BusinessStrategy #CompetitiveAdvantage #Innovation #MarketStrategy #Investing",
         filename="033_What_Is_a_Second_Mover_Advantage.docx", qc="PASS", notes="Pilot article, filename corrected from 003 to 033 (brief row 33). Brief tagged Indirect/Omit; a genuine connection (network effects/switching costs/patents as durability barriers) was found in the article's own body and added.")
    ,
    dict(no=71, seo="What Is MPLY ETF? Strategy Shares Monopoly ETF Explained",
         meta="MPLY is an actively managed ETF that screens for companies with monopoly-like traits. Here's what it holds, what it costs, and how it's built.",
         mply_decision="Primary Subject", mply_placement="Primary Subject",
         visual="Yes", visual_brief="Top-10 holdings bar chart, verified data (built).",
         internal_links="Anchor 'competitive moat' -> What Is a Competitive Moat in Investing (No.2); Anchor 'pricing power' -> What Is Pricing Power (No.1)",
         ext_sources="strategysharesetfs.com/mply; stockanalysis.com/etf/mply",
         hashtags="#MPLY #MPLYETF #ETFResearch #MonopolyInvesting #CompetitiveAdvantage #ActiveManagement #Investing",
         filename="071_What_Is_MPLY_ETF.docx", qc="PASS", notes="Pilot article, filename corrected from 004 to 071 (brief row 71).")
]

wb_t = openpyxl.load_workbook(TRACKER) if False else None
try:
    wb_t = openpyxl.load_workbook(TRACKER)
except FileNotFoundError:
    wb_t = None

if wb_t is None:
    # Start from the template.
    wb_t = openpyxl.load_workbook(
        r"C:\Users\SrimanNarayan\OneDrive - TransForm Solutions (P) Limited\MFunds_WP Articles\Production Trackers\_TEMPLATE_MPLY_Article_Production.xlsx"
    )
    ws = wb_t["Production Tracker"]
    # Clear any example rows below header (template has header only, but be safe).
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
else:
    ws = wb_t["Production Tracker"]

start_row = ws.max_row + 1 if ws.cell(row=2, column=1).value else 2

BODY_FONT = Font(name="Arial")
WRAP = Alignment(wrap_text=True, vertical="top")

for i, item in enumerate(PROD):
    r = start_row + i
    b = brief_rows[item["no"]]
    row_vals = [
        "2026-09-28",                       # Date
        "Set 01",                           # Set
        item["no"],                         # No.
        b["title"],                         # Article Title
        b["primary_kw"],                    # Primary Keyword
        b["secondary_kw"],                  # Secondary Keywords
        b["related"],                       # Related Terms
        b["cluster"],                       # Content Cluster
        b["intent"],                        # Search Intent
        b["relevance"],                     # MPLY Relevance
        b["connection_logic"],              # MPLY Connection Logic
        item["seo"],                        # SEO Title
        item["meta"],                       # Meta Description
        item["mply_decision"],              # MPLY Decision
        item["mply_placement"],             # MPLY Placement
        item["visual"],                     # Visual Decision
        item["visual_brief"],               # Visual Brief
        item["internal_links"],             # Internal Link Suggestions
        item["ext_sources"],                # External Sources
        item["hashtags"],                   # Social Hashtags
        item["filename"],                   # Article DOCX Filename
        item["qc"],                         # Final QC Status
        item["notes"],                      # Notes
    ]
    for c, v in enumerate(row_vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = BODY_FONT
        cell.alignment = WRAP
    ws.row_dimensions[r].height = 30

wb_t.save(TRACKER)
print(f"Wrote {len(PROD)} rows starting at row {start_row} to {TRACKER}")
